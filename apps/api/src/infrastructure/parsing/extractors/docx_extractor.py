import io

from docx import Document

from src.infrastructure.parsing.extractors.base import ExtractionResult, TextExtractor


class DocxExtractor(TextExtractor):
    def extract(self, content: bytes) -> ExtractionResult:
        doc = Document(io.BytesIO(content))
        paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
        raw_text = "\n".join(paragraphs).strip()

        if not raw_text:
            raise ValueError("No extractable text found in DOCX file.")

        return ExtractionResult(
            raw_text=raw_text,
            page_or_section_count=len(doc.paragraphs),
            word_count=len(raw_text.split()),
        )